from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.reports.matrix import build_table, is_total_label


def export_xlsx(
    path: Path,
    columns: list[str],
    rows: list[dict[str, Any]],
    totals: dict[str, Any],
    *,
    period_label: str | None = None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    values = build_table(columns, rows, totals, period_label=period_label)
    for row in values:
        sheet.append(list(row))
    if values:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row_index in range(2, sheet.max_row + 1):
            if is_total_label(sheet.cell(row_index, 1).value):
                for cell in sheet[row_index]:
                    cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    if columns and values:
        last_data_row = len(values)
        while (
            last_data_row > 1
            and values[last_data_row - 1]
            and is_total_label(values[last_data_row - 1][0])
        ):
            last_data_row -= 1
        column_count = max(len(row) for row in values)
        if column_count and last_data_row >= 1:
            sheet.auto_filter.ref = f"A1:{get_column_letter(column_count)}{max(last_data_row, 1)}"
    if values:
        column_count = max((len(row) for row in values), default=0)
        for index in range(1, column_count + 1):
            widths = [
                str(row[index - 1]) if index - 1 < len(row) and row[index - 1] is not None else ""
                for row in values[:200]
            ]
            sheet.column_dimensions[get_column_letter(index)].width = min(
                max((len(value) for value in widths), default=0) + 2,
                45,
            )
    workbook.save(path)
