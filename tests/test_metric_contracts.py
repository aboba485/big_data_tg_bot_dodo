from __future__ import annotations

import csv
from datetime import date
from typing import Any

import httpx
import pytest
from openpyxl import load_workbook

from app.dodo.profile import ProfileResolver
from app.errors import ConfigurationError
from app.metric_contracts import validate_metric_contracts
from app.planner.schemas import Granularity, PlanMode, PlanStatus, ReportPlan
from app.reports.exporters import export_csv, export_xlsx
from app.reports.metric_registry import MetricDefinition, MetricRegistry
from app.reports.service import ReportAggregator
from app.services import build_services
from tests.conftest import UNIT_ID

SPECIAL_GROUP_COLUMNS = {
    "sales channel": "salesChannel",
    "ingredient": "ingredient",
    "ingredient category": "ingredientCategory",
    "stop reason": "stopReason",
}


def _plan(
    metric_id: str,
    registry: MetricRegistry,
    *,
    granularity: Granularity = Granularity.TOTAL,
    special_groups: list[str] | None = None,
) -> ReportPlan:
    definition = registry.get(metric_id)
    group_by = ["unit", *(special_groups or [])]
    if granularity != Granularity.TOTAL:
        group_by.insert(0, granularity.value)
    return ReportPlan(
        status=PlanStatus.READY,
        mode=PlanMode.METRICS,
        metric_ids=[metric_id],
        operation_ids=[definition.operation_id],
        date_from=date(2026, 6, 1),
        date_to=date(2026, 6, 2),
        granularity=granularity,
        group_by=group_by,
    )


class EmptyMetricClient:
    def __init__(self, profiles: ProfileResolver, registry: MetricRegistry) -> None:
        self.profiles = profiles
        self.registry = registry

    async def request_operation(
        self,
        operation: Any,
        _query_params: dict[str, Any],
        _path_values: dict[str, str] | None = None,
    ) -> dict[str, Any]:
        profile = self.profiles.resolve(operation)
        if profile.pagination:
            return {profile.items_path: [], profile.end_flag_path: True}
        if profile.collection:
            return {profile.collection: []}
        return {
            field: 0
            for metric_id in self.registry.all()
            if self.registry.get(metric_id).operation_id == operation.operation_id
            for field in self.registry.get(metric_id).required_fields
            if "." not in field
        }


@pytest.mark.asyncio
async def test_every_registered_metric_executes_aggregates_and_exports(settings, tmp_path) -> None:
    async with httpx.AsyncClient() as http_client:
        services = build_services(settings, http_client)
        orchestrator = services["orchestrator"]
        registry = services["metrics"]
        for metric_id in sorted(registry.all()):
            definition = registry.get(metric_id)
            default_groups = (
                ["sales channel"] if definition.nested_collection == "salesBreakdown" else []
            )
            plan = _plan(metric_id, registry, special_groups=default_groups)
            records = await orchestrator.executor.execute(plan, [UNIT_ID])
            report = orchestrator.aggregator.aggregate(plan, records)

            assert report.rows, metric_id
            assert metric_id in report.columns, metric_id
            assert report.totals.get(metric_id) is not None, metric_id

            csv_path = tmp_path / f"{metric_id}.csv"
            xlsx_path = tmp_path / f"{metric_id}.xlsx"
            export_csv(csv_path, report.columns, report.rows, report.totals)
            export_xlsx(xlsx_path, report.columns, report.rows, report.totals)
            assert csv_path.stat().st_size > 0, metric_id
            assert xlsx_path.stat().st_size > 0, metric_id
            with csv_path.open(encoding="utf-8-sig", newline="") as stream:
                csv_values = list(csv.reader(stream))
            assert any(metric_id in row for row in csv_values), metric_id
            workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
            try:
                xlsx_values = {
                    str(cell)
                    for row in workbook.active.iter_rows(values_only=True)
                    for cell in row
                    if cell is not None
                }
            finally:
                workbook.close()
            assert metric_id in xlsx_values, metric_id

            for granularity_name in definition.granularities:
                granularity = Granularity(granularity_name)
                granularity_plan = _plan(
                    metric_id,
                    registry,
                    granularity=granularity,
                    special_groups=default_groups,
                )
                granular_records = await orchestrator.executor.execute(granularity_plan, [UNIT_ID])
                granular_report = orchestrator.aggregator.aggregate(
                    granularity_plan, granular_records
                )
                assert granular_report.rows, f"{metric_id}:{granularity_name}"
                if granularity != Granularity.TOTAL:
                    assert granularity.value in granular_report.columns
                    assert all(
                        row.get(granularity.value) is not None for row in granular_report.rows
                    )

            for group, column in SPECIAL_GROUP_COLUMNS.items():
                if group not in definition.groups:
                    continue
                group_plan = _plan(metric_id, registry, special_groups=[group])
                group_records = await orchestrator.executor.execute(group_plan, [UNIT_ID])
                group_report = orchestrator.aggregator.aggregate(group_plan, group_records)
                assert column in group_report.columns, f"{metric_id}:{group}"
                assert all(row.get(column) is not None for row in group_report.rows)


@pytest.mark.asyncio
async def test_every_registered_metric_handles_empty_dodo_data(settings) -> None:
    async with httpx.AsyncClient() as http_client:
        services = build_services(settings, http_client)
        orchestrator = services["orchestrator"]
        registry = services["metrics"]
        orchestrator.executor.client = EmptyMetricClient(orchestrator.executor.profiles, registry)

        for metric_id in sorted(registry.all()):
            plan = _plan(metric_id, registry)
            records = await orchestrator.executor.execute(plan, [UNIT_ID])
            report = orchestrator.aggregator.aggregate(plan, records)

            assert metric_id in report.totals, metric_id
            assert metric_id in report.columns, metric_id


def test_every_registered_metric_uses_its_formula_fields() -> None:
    registry = MetricRegistry()
    aggregator = ReportAggregator(registry)
    for metric_id in sorted(registry.all()):
        definition = registry.get(metric_id)
        rows = [
            {
                "unitId": UNIT_ID,
                "date": "2026-06-01",
                "__bucket_start": "2026-06-01",
                "__bucket_end": "2026-06-01",
            },
            {
                "unitId": UNIT_ID,
                "date": "2026-06-02",
                "__bucket_start": "2026-06-02",
                "__bucket_end": "2026-06-02",
            },
        ]
        if definition.aggregation == "sum":
            rows[0][str(definition.value_field)] = 10
            rows[1][str(definition.value_field)] = 20
            expected = 30.0
        elif definition.aggregation == "count_unique":
            rows[0][str(definition.value_field)] = "a"
            rows[1][str(definition.value_field)] = "b"
            expected = 2
        elif definition.aggregation == "weighted_average":
            rows[0][str(definition.value_field)] = 10
            rows[1][str(definition.value_field)] = 20
            rows[0][str(definition.weight_field)] = 1
            rows[1][str(definition.weight_field)] = 3
            expected = 17.5
        elif definition.aggregation == "ratio_from_sums":
            rows[0][str(definition.numerator_field)] = 20
            rows[1][str(definition.numerator_field)] = 30
            rows[0][str(definition.denominator_field)] = 4
            rows[1][str(definition.denominator_field)] = 6
            expected = 5.0 * definition.multiplier
        elif definition.aggregation == "duration_sum":
            rows[0].update(
                id="a",
                startedAtLocal="2026-06-01T00:00:00",
                endedAtLocal="2026-06-01T01:00:00",
            )
            rows[1].update(
                id="b",
                startedAtLocal="2026-06-02T02:00:00",
                endedAtLocal="2026-06-02T03:00:00",
            )
            expected = 2.0
        else:
            pytest.fail(f"Нет независимого oracle для {metric_id}:{definition.aggregation}")

        plan = _plan(metric_id, registry)
        report = aggregator.aggregate(plan, {definition.operation_id: rows})

        assert report.totals[metric_id] == pytest.approx(expected), metric_id


@pytest.mark.asyncio
async def test_metric_contract_validator_reports_documentation_drift(settings) -> None:
    async with httpx.AsyncClient() as http_client:
        services = build_services(settings, http_client)
        registry = MetricRegistry()
        registry.get("sales").required_fields.append("missingFromDocumentation")

        with pytest.raises(ConfigurationError, match="missingFromDocumentation"):
            validate_metric_contracts(
                registry,
                services["repository"],
                settings.allowed_operations,
                ProfileResolver(settings.endpoint_overrides),
            )


@pytest.mark.asyncio
async def test_metric_contract_validator_rejects_disabled_documented_pagination(settings) -> None:
    async with httpx.AsyncClient() as http_client:
        services = build_services(settings, http_client)
        overrides = {
            **settings.endpoint_overrides,
            "get-delivery-vouchers": {
                **settings.endpoint_overrides["get-delivery-vouchers"],
                "pagination": False,
            },
        }

        with pytest.raises(ConfigurationError, match="пагинация из документации отключена"):
            validate_metric_contracts(
                MetricRegistry(),
                services["repository"],
                settings.allowed_operations,
                ProfileResolver(overrides),
            )


@pytest.mark.parametrize(
    "definition",
    [
        {
            "operation_id": "example",
            "collection": "items",
            "aggregation": "typo",
            "required_fields": ["value"],
            "granularities": ["total"],
            "groups": ["unit"],
        },
        {
            "operation_id": "example",
            "collection": "items",
            "aggregation": "weighted_average",
            "value_field": "value",
            "required_fields": ["value"],
            "granularities": ["total"],
            "groups": ["unit"],
        },
        {
            "operation_id": "example",
            "collection": "items",
            "aggregation": "duration_sum",
            "required_fields": ["startedAtLocal", "endedAtLocal"],
            "granularities": ["total"],
            "groups": ["unit"],
        },
    ],
)
def test_metric_definition_rejects_incomplete_formula(definition) -> None:
    with pytest.raises(ValueError):
        MetricDefinition.model_validate(definition)
