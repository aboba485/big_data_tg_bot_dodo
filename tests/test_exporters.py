from __future__ import annotations

from openpyxl import load_workbook

from app.reports.exporters import export_csv, export_xlsx


def test_xlsx_unit_date_matrix_has_no_metadata(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    export_xlsx(
        path,
        ["day", "unitName", "sales"],
        [
            {"day": "2026-06-01", "unitName": "Москва 4-3", "sales": 235972},
            {"day": "2026-06-02", "unitName": "Москва 4-3", "sales": 215717},
        ],
        {"sales": 451689},
    )

    workbook = load_workbook(path)
    sheet = workbook.active
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    dumped = str(values)
    assert "Отчёт Dodo IS" not in dumped
    assert "Запрос" not in dumped
    assert values[0][0] in {"", None}
    assert values[0][1:] == ["01.06.2026", "02.06.2026"]
    assert values[1] == ["Москва 4-3", 235972, 215717]
    assert values[2] == ["Итого", 235972, 215717]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:C2"


def test_xlsx_without_units_stays_long_and_filters_all_data_rows(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    export_xlsx(
        path,
        ["day", "sales"],
        [
            {"day": "2026-06-01", "sales": 100},
            {"day": "2026-06-02", "sales": 200},
        ],
        {},
    )

    workbook = load_workbook(path)
    sheet = workbook.active
    assert [cell.value for cell in sheet[1]] == ["day", "sales"]
    assert sheet.auto_filter.ref == "A1:B3"


def test_xlsx_empty_columns_freeze_without_autofilter(tmp_path) -> None:
    path = tmp_path / "report.xlsx"
    export_xlsx(path, [], [], {})

    workbook = load_workbook(path)
    sheet = workbook.active
    assert sheet.freeze_panes == "A2"
    assert not sheet.auto_filter.ref


def test_csv_matrix_pivots_units_across_dates(tmp_path) -> None:
    path = tmp_path / "report.csv"
    export_csv(
        path,
        ["day", "unitName", "sales"],
        [
            {"day": "2026-06-01", "unitName": "Москва 4-3", "sales": 235972},
            {"day": "2026-06-02", "unitName": "Москва 4-3", "sales": 215717},
        ],
        {"sales": 451689},
    )

    text = path.read_text(encoding="utf-8-sig")
    assert ",01.06.2026,02.06.2026" in text
    assert "Москва 4-3,235972,215717" in text
    assert "Итого,235972,215717" in text
